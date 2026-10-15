import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from data import DataManager
from http_client import AuthenticatedHttpClient
from office import OfficeDocumentManager
from template import TemplateManager


class HttpClientRegressionTests(unittest.TestCase):
    def test_authenticated_client_applies_shared_header_and_timeout(self):
        client = AuthenticatedHttpClient(timeout=17)
        response = object()

        with patch.object(client.session, "get", return_value=response) as get:
            self.assertIs(
                client.get("https://example.com/resource", auth_key="token"),
                response,
            )

        get.assert_called_once_with(
            "https://example.com/resource",
            headers={"authorization": "Bearer token"},
            timeout=17,
        )

    def test_request_log_contains_content_without_credentials(self):
        client = AuthenticatedHttpClient()
        with patch.object(client.session, "get", return_value=object()) as get, patch(
            "http_client.logger"
        ) as logger:
            client.get(
                "https://example.com/resource",
                auth_key="secret-token",
                headers={
                    "Authorization": "stale-token",
                    "X-Trace": "trace-id",
                    "Cookie": "session=secret",
                },
                params={"project": "123"},
                json={"enabled": True},
            )

        message = logger.info.call_args.args
        self.assertIn("HTTP request method=GET", message[0])
        self.assertNotIn("secret-token", str(message))
        self.assertNotIn("session=secret", str(message))
        self.assertIn("trace-id", str(message))
        self.assertIn("project", str(message))
        self.assertIn("enabled", str(message))
        get.assert_called_once_with(
            "https://example.com/resource",
            headers={
                "authorization": "Bearer " + "secret-token",
                "X-Trace": "trace-id",
                "Cookie": "session=secret",
            },
            timeout=30,
            params={"project": "123"},
            json={"enabled": True},
        )


class DataNormalizationTests(unittest.TestCase):
    def test_missing_battery_data_classifies_system_as_pv(self):
        data = DataManager()

        data.load_json(
            {
                "id": 123,
                "public_id": "project-api-id",
                "client_name": "Ada Lovelace",
                "Materials": {},
            },
            save_to_file=False,
        )

        self.assertEqual(data.system.type, 0)

    def test_load_json_normalizes_project_customer_system_and_consumption(self):
        data = DataManager()
        payload = {
            "id": 123,
            "public_id": "project-api-id",
            "client_name": "Ada Lovelace",
            "client_phone": "555-0100",
            "Position": {
                "street": "1 Analytical Engine Way",
                "city": "London",
                "region_code": "CA",
                "postal_code": "90210",
            },
            "is_roofing_only": True,
            "proposals": [{"sizeInKw": 8.4}],
            "Materials": {
                "panel": [
                    {
                        "Manufacturer": {"name": "Acme Solar"},
                        "name": "Acme 400",
                        "type": "mono",
                        "size_in_watts": 400,
                        "warranty": 25,
                        "ptc": 380,
                        "count": 21,
                        "spec_sheet": "https://example.com/panel.pdf",
                    }
                ],
                "inverter": [
                    {
                        "Manufacturer": {"name": "Acme Inverters"},
                        "name": "Acme 8K",
                        "warranty": 12,
                        "count": 1,
                        "spec_sheet": "https://example.com/inverter.pdf",
                    }
                ],
            },
            "Settings": {
                "Pricing": {
                    "ProjectConnection": {
                        "genability_lse_name": "Example Utility",
                        "proposal_id": "proposal-id",
                        "Bills": [
                            {"consumptionInKwh": 400},
                            {"consumptionInKwh": 650.5},
                        ],
                    },
                    "FinancialOptions": [{"id": "financial-id"}],
                }
            },
        }

        data.load_json(payload, save_to_file=False)

        self.assertEqual(data.project.id, 123)
        self.assertEqual(data.client.first_name, "Ada")
        self.assertEqual(data.client.last_name, "Lovelace")
        self.assertEqual(data.system.type, 0)
        self.assertEqual(data.system.pv_size, 8.4)
        self.assertEqual(data.system.panel.size, 400)
        self.assertEqual(data.utility.annual_consumption, 1050.5)
        self.assertEqual(data.project.financial_id, "financial-id")
        self.assertEqual(data.project.proposal_id, "proposal-id")


class TemplateRegressionTests(unittest.TestCase):
    def test_pv_templates_do_not_access_battery_fields(self):
        data = DataManager()
        data.system.type = 0
        data.project.name = "Ada Lovelace"
        data.client.name = "Ada Lovelace"
        data.utility.name = "Example Utility"
        data.system.panel.type = "mono"
        data.system.inverter.manufacturer = "Acme Inverters"
        data.pricing.loan_interest_rate = 0

        class BatteryAccessGuard:
            def __getattribute__(self, name):
                raise AssertionError(f"PV template accessed battery field: {name}")

        data.system.battery = BatteryAccessGuard()

        templates = TemplateManager(data)

        self.assertNotIn("{{BATTERY_MANUFACTURER}}", templates.get_all())
        self.assertNotIn("{{BATTERY_PLURAL}}", templates.get_all())

    def test_settings_aliases_and_derived_values_are_applied(self):
        with tempfile.TemporaryDirectory() as directory:
            settings_path = Path(directory) / "settings.xlsx"
            workbook = Workbook()
            aliases = workbook.active
            aliases.title = "Aliases"
            aliases.append(["placeholder", "replacement"])
            aliases.append(["Example Utility", "Example Power"])
            variables = workbook.create_sheet("Variables")
            variables.append(["placeholder", "value"])
            variables.append(["{{UTILITY_NEM_PGE}}", "PGE NEM terms"])
            variables.append(["{{UTILITY_APPLICATION_DESCRIPTION_PGE}}", "PGE application"])
            variables.append(["{{UTILITY_LINK_PGE}}", "https://example.com"])
            variables.append(["{{APPLICATION_FEE_PGE}}", "$0"])
            workbook.create_sheet("Spec Sheets")
            workbook.create_sheet("Spreadsheet")
            workbook.save(settings_path)

            data = DataManager()
            data.files.settings = settings_path
            data.project.name = "Ada Lovelace"
            data.client.name = "Ada Lovelace"
            data.client.first_name = "Ada"
            data.client.last_name = "Lovelace"
            data.utility.name = "Example Utility"
            data.utility.annual_consumption = 1050.5
            data.system.type = 0
            data.system.pv_size = 8.4
            data.system.panel.manufacturer = "Acme Solar"
            data.system.panel.model = "Acme 400"
            data.system.panel.type = "mono"
            data.system.panel.size = 400
            data.system.panel.ptc = 380
            data.system.panel.count = 21
            data.system.inverter.manufacturer = "Acme Inverters"
            data.system.inverter.model = "Acme 8K"
            data.system.inverter.count = 1
            data.pricing.pv_cost = 24000
            data.pricing.total_cost = 24000
            data.pricing.loan_interest_rate = 0.05
            data.pricing.loan_monthly_payment = 500

            templates = TemplateManager(data)

            self.assertEqual(templates.resolve("{{UTILITY_NAME}}"), "Example Power")
            self.assertEqual(templates.resolve("{{UTILITY_NEM}}"), "PGE NEM terms")
            self.assertEqual(templates.resolve("{{PANEL_PTC_RATIO}}"), 95)
            self.assertEqual(templates.resolve("{{LOAN_INTEREST_RATE}}"), "5.00")
            self.assertEqual(templates.resolve("{{TOTAL_COST}}"), "$24,000")


class FinalizationRegressionTests(unittest.TestCase):
    def test_finalize_combines_cover_letter_proposal_and_spec_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            documents = root / "documents"
            specsheets = documents / "specsheets"
            specsheets.mkdir(parents=True)

            settings_path = root / "settings.xlsx"
            workbook = Workbook()
            workbook.active.title = "Aliases"
            workbook.create_sheet("Variables")
            spec_sheets = workbook.create_sheet("Spec Sheets")
            spec_sheets.append(["model", "pages"])
            spec_sheets.append(["Acme 400", 1])
            spec_sheets.append(["Acme 8K", 1])
            workbook.create_sheet("Spreadsheet")
            workbook.save(settings_path)

            cover_template = root / "cover_letter.docx"
            Document().save(cover_template)
            cover_letter_docx = documents / "cover_letter.docx"
            Document().save(cover_letter_docx)

            proposal_pdf = documents / "solargraf_proposal.pdf"
            self._write_pdf(proposal_pdf, 1)
            pricing_spreadsheet = documents / "pricing_spreadsheet.xlsx"
            Workbook().save(pricing_spreadsheet)

            data = DataManager()
            data.project.id = 123
            data.project.api_id = "api-id"
            data.project.proposal_id = "proposal-id"
            data.project.financial_id = "financial-id"
            data.client.name = "Ada Lovelace"
            data.client.first_name = "Ada"
            data.client.last_name = "Lovelace"
            data.client.address.street = "1 Analytical Engine Way"
            data.client.address.city = "London"
            data.client.address.region_code = "CA"
            data.client.address.postal_code = "90210"
            data.utility.name = "Example Utility"
            data.system.type = 0
            data.system.pv_size = 8.4
            data.system.panel.manufacturer = "Acme Solar"
            data.system.panel.model = "Acme 400"
            data.system.panel.size = 400
            data.system.panel.count = 21
            data.system.panel.specsheet_url = "https://example.com/panel.pdf"
            data.system.inverter.manufacturer = "Acme Inverters"
            data.system.inverter.model = "Acme 8K"
            data.system.inverter.count = 1
            data.system.inverter.specsheet_url = "https://example.com/inverter.pdf"
            data.files.settings = settings_path
            data.files.cover_letter_template = cover_template
            data.files.cover_letter_docx = cover_letter_docx
            data.files.cover_letter_pdf = documents / "cover_letter.pdf"
            data.files.solargraf_proposal = proposal_pdf
            data.files.pricing_spreadsheet = pricing_spreadsheet
            data.files.specsheets_folder = specsheets
            data.files.final_proposal = documents / "final.pdf"

            class Variables:
                def resolve(self, name):
                    return {
                        "{{PANEL_MODEL}}": "Acme 400",
                        "{{INVERTER_MODEL}}": "Acme 8K",
                    }.get(name, name)

            class Scraper:
                def download(self, url, output_dir):
                    output = output_dir / Path(url).name
                    self._write_pdf(output, 1)
                    return output

                @staticmethod
                def _write_pdf(path, pages):
                    writer = PdfWriter()
                    for _ in range(pages):
                        writer.add_blank_page(width=612, height=792)
                    with path.open("wb") as stream:
                        writer.write(stream)

            def convert_cover_letter(source, target):
                self._write_pdf(Path(target), 1)

            with patch(
                "office.docx2pdf.convert", side_effect=convert_cover_letter
            ) as convert:
                result = OfficeDocumentManager().finalize_proposal(
                    data, Variables(), Scraper()
                )

            self.assertEqual(result, data.files.final_proposal)
            self.assertTrue(result.is_file())
            self.assertEqual(len(PdfReader(str(result)).pages), 4)
            convert.assert_called_once_with(
                str(cover_letter_docx), str(data.files.cover_letter_pdf)
            )

    @staticmethod
    def _write_pdf(path, pages):
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=612, height=792)
        with path.open("wb") as stream:
            writer.write(stream)


if __name__ == "__main__":
    unittest.main()
